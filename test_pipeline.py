"""End-to-end checks: build a sample password-protected statement, convert it, verify."""

from __future__ import annotations

import io
import sys

import pymupdf as fitz
import pandas as pd
from openpyxl import load_workbook

from excel_writer import build_workbook
from mpesa_pdf import ExtractionError, PasswordError, extract_statement

PASSWORD = "S3cret123"

HEADER = [
    "Receipt No.", "Completion Time", "Details", "Transaction Status",
    "Paid In", "Withdrawn", "Balance",
]

ROWS = [
    ["RFK01ABCDEF", "01/07/2025 08:15:22", "PAYMENT TO 254712345678 JOHN DOE", "Completed", "", "1500.00", "8500.00"],
    ["RFK02GHIJKL", "02/07/2025 11:03:45", "PAYMENT RECEIVED FROM JANE WAMBUI", "Completed", "20000.00", "", "28500.00"],
    ["SFH03MNOPQR", "05/07/2025 17:44:10", "ATM WITHDRAWAL AT KAREN", "Completed", "", "5000.00", "23500.00"],
    ["SBL04STUVWX", "10/07/2025 09:20:00", "LIPA NA M-PESA PAYBILL 400200", "Completed", "", "1250.75", "22249.25"],
    ["RFK05YZABCD", "15/07/2025 13:05:59", "PAYMENT RECEIVED FROM ACME LTD", "Completed", "7500.50", "", "29749.75"],
]


def build_sample_pdf(password: str = PASSWORD) -> bytes:
    """Create a ruled-table M-PESA-style statement and encrypt it."""
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)  # landscape A4
    page.insert_text((40, 45), "M-PESA STATEMENT", fontsize=16)
    page.insert_text((40, 65), "Name: JANE WAMBUI   M-PESA Number: 254712345678", fontsize=10)

    x0, y0 = 40.0, 90.0
    widths = [80, 118, 275, 82, 72, 78, 75]  # sums to 780; fits in 842 - 40 margin
    assert sum(widths) <= 842 - x0 - 10, "columns overflow the page"
    row_h = 26.0
    rows = [HEADER] + ROWS

    for r, row in enumerate(rows):
        x = x0
        for c, text in enumerate(row):
            rect = fitz.Rect(x, y0 + r * row_h, x + widths[c], y0 + (r + 1) * row_h)
            page.draw_rect(rect, color=(0, 0, 0), width=0.6)
            align = fitz.TEXT_ALIGN_RIGHT if c >= 4 else fitz.TEXT_ALIGN_LEFT
            inner = rect + (4, 6, -4, -4)
            if page.insert_textbox(inner, str(text), fontsize=8, align=align) < 0:
                raise RuntimeError(f"text did not fit cell: {text!r}")
            x += widths[c]

    payload = doc.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        user_pw=password,
        owner_pw=password,
    )
    doc.close()
    return payload


UNRULED_COLS = [40, 130, 240, 430, 500, 560, 630]
UNRULED_LABELS = [
    "Receipt No.", "Completion Time", "Details", "Transaction Status",
    "Paid In", "Withdrawn", "Balance",
]
UNRULED_ROWS = [
    ["RFK99XYZ", "01/08/2025 09:00:00", "PAYMENT TO 254700000000", "Completed", "1000.00", "", "9000.00"],
    ["RFK98ABC", "02/08/2025 10:30:00", "ATM WITHDRAWAL UGB", "Completed", "", "500.00", "8500.00"],
]


def build_unruled_pdf() -> bytes:
    """A statement whose columns are implied by position only (no table lines)."""
    doc = fitz.open()
    page = doc.new_page(width=760, height=220)
    page.insert_text((40, 50), "M-PESA STATEMENT (no ruling lines)", fontsize=13)
    for x, label in zip(UNRULED_COLS, UNRULED_LABELS):
        page.insert_text((x, 90), label, fontsize=8)
    for i, row in enumerate(UNRULED_ROWS):
        for x, value in zip(UNRULED_COLS, row):
            if value:
                page.insert_text((x, 120 + i * 20), value, fontsize=8)
    payload = doc.tobytes()
    doc.close()
    return payload


def check(label: str, condition: bool, detail: str = "") -> bool:
    status = "PASS" if condition else "FAIL"
    print(f"[{status}] {label}" + (f" -- {detail}" if detail and not condition else ""))
    return condition


def main() -> int:
    pdf_bytes = build_sample_pdf()
    ok = True

    # 1. Wrong password is rejected
    try:
        extract_statement(pdf_bytes, "wrong-password")
        ok &= check("wrong password rejected", False)
    except PasswordError:
        ok &= check("wrong password rejected", True)

    # 2. Missing password is rejected
    try:
        extract_statement(pdf_bytes, "")
        ok &= check("missing password rejected", False)
    except PasswordError:
        ok &= check("missing password rejected", True)

    # 3. Correct password yields all rows
    statement = extract_statement(pdf_bytes, PASSWORD)
    frame = statement.transactions
    ok &= check("row count", len(frame) == len(ROWS), f"got {len(frame)}")
    ok &= check("columns", list(frame.columns) == [
        "receipt_no", "completion_time", "details", "status",
        "paid_in", "withdrawn", "balance",
    ], str(list(frame.columns)))
    ok &= check("amount parsed", frame["paid_in"].iloc[1] == 20000.0,
                str(frame["paid_in"].iloc[1]))
    ok &= check("blank amount is null", pd.isna(frame["withdrawn"].iloc[1]))
    ok &= check("datetime parsed",
                str(frame["completion_time"].iloc[0]) == "2025-07-01 08:15:22",
                str(frame["completion_time"].iloc[0]))
    ok &= check("sorted by time",
                frame["completion_time"].is_monotonic_increasing)

    # 4. Excel output
    workbook_bytes = build_workbook(frame, "sample.pdf")
    ok &= check("xlsx magic", workbook_bytes[:2] == b"PK")
    wb = load_workbook(io.BytesIO(workbook_bytes))
    ok &= check("sheets", wb.sheetnames == ["Summary", "Transactions"], str(wb.sheetnames))
    ws = wb["Transactions"]
    ok &= check("header row", ws.cell(1, 1).value == "Receipt No.")
    ok &= check("data row count", ws.max_row == len(ROWS) + 1, str(ws.max_row))
    ok &= check("number format", ws.cell(2, 5).number_format.startswith("#,##0.00"),
                str(ws.cell(2, 5).number_format))
    ok &= check("summary total", wb["Summary"].cell(9, 2).value is not None)

    # 5. Non-PDF input
    try:
        extract_statement(b"not a pdf at all", "x")
        ok &= check("non-PDF rejected", False)
    except ExtractionError:
        ok &= check("non-PDF rejected", True)

    # 5b. Statement with no ruling lines (text-coordinate fallback)
    plain = build_unruled_pdf()
    unruled = extract_statement(plain, "")
    uf = unruled.transactions
    ok &= check("unruled row count", len(uf) == 2, str(len(uf)))
    ok &= check("unruled columns", list(uf.columns) == [
        "receipt_no", "completion_time", "details", "status",
        "paid_in", "withdrawn", "balance",
    ], str(list(uf.columns)))
    ok &= check("unruled amount", uf["paid_in"].iloc[0] == 1000.0, str(uf["paid_in"].iloc[0]))
    ok &= check("unruled date",
                str(uf["completion_time"].iloc[1]) == "2025-08-02 10:30:00",
                str(uf["completion_time"].iloc[1]))

    # 5c. Multi-page statements, with and without a repeated header row
    import test_multipage as mp

    for flag in (True, False):
        multi = extract_statement(mp.build_multipage_pdf(flag), "pw123")
        expected = len(mp.PAGE1) + len(mp.PAGE2)
        label = f"multipage header={flag}"
        ok &= check(f"{label} row count", len(multi.transactions) == expected,
                    str(len(multi.transactions)))
        ok &= check(f"{label} page count", multi.pages == 2, str(multi.pages))
        ok &= check(f"{label} last balance",
                    multi.transactions["balance"].iloc[-1] == 1750.0,
                    str(multi.transactions["balance"].iloc[-1]))

    # 6. Flask endpoint round-trip
    from app import app

    client = app.test_client()
    landing = client.get("/")
    ok &= check("GET /", landing.status_code == 200)

    bad = client.post("/convert", data={"statement": (io.BytesIO(pdf_bytes), "s.pdf"),
                                        "password": "nope"},
                      content_type="multipart/form-data")
    ok &= check("POST wrong password redirects", bad.status_code == 302)

    good = client.post("/convert", data={"statement": (io.BytesIO(pdf_bytes), "s.pdf"),
                                         "password": PASSWORD},
                       content_type="multipart/form-data")
    ok &= check("POST success", good.status_code == 200)
    ok &= check("download header",
                "attachment" in good.headers.get("Content-Disposition", ""),
                good.headers.get("Content-Disposition", ""))
    ok &= check("download is xlsx", good.data[:2] == b"PK")

    nonpdf = client.post("/convert", data={"statement": (io.BytesIO(b"hello"), "s.txt"),
                                           "password": ""},
                         content_type="multipart/form-data")
    ok &= check("non-PDF upload rejected", nonpdf.status_code == 302)

    print("\nALL PASSED" if ok else "\nSOME CHECKS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
