"""Write extracted M-PESA transactions into a formatted Excel workbook."""

from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1B5E20")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL = PatternFill("solid", fgColor="E8F5E9")
THIN = Side(style="thin", color="C8E6C9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

AMOUNT_COLUMNS = ("paid_in", "withdrawn", "balance", "transaction_cost", "balance_after")
DATE_FORMAT = "dd/mm/yyyy hh:mm:ss"
MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'


def _header_label(column: str) -> str:
    label = column.replace("_", " ").title()
    return f"{label}." if label.endswith(" No") else label


def build_workbook(frame: pd.DataFrame, source_name: str = "statement.pdf") -> bytes:
    """Return the bytes of a formatted .xlsx workbook for the given transactions."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    headers = list(frame.columns)
    ws.append([_header_label(col) for col in headers])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 24

    for row_index, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        ws.append([_excel_value(col, val) for col, val in zip(headers, row)])
        for col_index, column in enumerate(headers, start=1):
            cell = ws.cell(row=row_index, column=col_index)
            cell.border = BORDER
            if column == "completion_time":
                cell.number_format = DATE_FORMAT
                cell.alignment = Alignment(horizontal="center")
            elif column in AMOUNT_COLUMNS:
                cell.number_format = MONEY_FORMAT
            elif column == "details":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_index % 2 == 0:
                cell.fill = ALT_FILL

    for col_index, column in enumerate(headers, start=1):
        letter = get_column_letter(col_index)
        longest = max(
            [len(_header_label(column))]
            + [len(str(v)) for v in frame[column].head(200) if v is not None]
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    _add_summary(wb, frame, source_name)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _excel_value(column: str, value: object) -> object:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if column == "completion_time":
        return pd.Timestamp(value).to_pydatetime()
    return value


def _total(frame: pd.DataFrame, column: str) -> float:
    """Sum a money column, tolerating columns the statement did not have."""
    if column not in frame.columns:
        return 0.0
    total = pd.to_numeric(frame[column], errors="coerce").sum()
    return float(total) if pd.notna(total) else 0.0


def _add_summary(wb: Workbook, frame: pd.DataFrame, source_name: str) -> None:
    """Add a Summary sheet with totals and statement info."""
    summary = wb.create_sheet("Summary", 0)
    paid = _total(frame, "paid_in")
    withdrawn = _total(frame, "withdrawn")
    costs = _total(frame, "transaction_cost")

    rows = [
        ("M-PESA Statement Summary", ""),
        ("Source file", source_name),
        ("Generated", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("Transactions", int(len(frame))),
        ("Period start", _min_time(frame)),
        ("Period end", _max_time(frame)),
        ("Total received (Paid In)", float(paid)),
        ("Total withdrawn", float(withdrawn)),
        ("Total transaction costs", float(costs)),
        ("Net cash flow", float(paid - withdrawn)),
    ]

    for row in rows:
        summary.append(list(row))

    summary["A1"].font = Font(bold=True, size=14, color="1B5E20")
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    for row in summary.iter_rows(min_row=7, max_row=10, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = MONEY_FORMAT
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 42


def _min_time(frame: pd.DataFrame) -> str:
    if frame.empty or "completion_time" not in frame:
        return ""
    return pd.Timestamp(frame["completion_time"].min()).strftime("%d/%m/%Y %H:%M")


def _max_time(frame: pd.DataFrame) -> str:
    if frame.empty or "completion_time" not in frame:
        return ""
    return pd.Timestamp(frame["completion_time"].max()).strftime("%d/%m/%Y %H:%M")
