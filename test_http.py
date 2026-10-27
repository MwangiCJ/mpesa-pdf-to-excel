"""Live HTTP checks against a running server (python app.py)."""

from __future__ import annotations

import io
import sys

import requests
from openpyxl import load_workbook

import test_pipeline as t

BASE = "http://127.0.0.1:5000"


def check(label: str, condition: bool, detail: str = "") -> bool:
    print(f"[{'PASS' if condition else 'FAIL'}] {label}" + (f" -- {detail}" if detail and not condition else ""))
    return condition


def main() -> int:
    pdf = t.build_sample_pdf()
    ok = True

    r = requests.get(BASE + "/")
    ok &= check("GET /", r.status_code == 200 and 'name="password"' in r.text)

    r = requests.post(
        BASE + "/convert",
        files={"statement": ("stmt.pdf", pdf, "application/pdf")},
        data={"password": "nope"},
        allow_redirects=False,
    )
    ok &= check("wrong password redirects", r.status_code == 302, str(r.status_code))

    r = requests.post(
        BASE + "/convert",
        files={"statement": ("stmt.pdf", pdf, "application/pdf")},
        data={"password": t.PASSWORD},
    )
    disposition = r.headers.get("Content-Disposition", "")
    ok &= check("convert returns 200", r.status_code == 200, str(r.status_code))
    ok &= check(
        "attachment header",
        "attachment" in disposition and disposition.split("=", 1)[-1].endswith("_transactions.xlsx"),
        disposition,
    )
    ok &= check("body is xlsx", r.content[:2] == b"PK", str(r.content[:8]))
    ok &= check("output name", "_transactions.xlsx" in disposition, disposition)

    wb = load_workbook(io.BytesIO(r.content))
    ok &= check("sheets", wb.sheetnames == ["Summary", "Transactions"], str(wb.sheetnames))
    ws = wb["Transactions"]
    ok &= check("5 data rows", ws.max_row == 6, str(ws.max_row))
    ok &= check("first receipt", ws.cell(2, 1).value == "RFK01ABCDEF", str(ws.cell(2, 1).value))
    ok &= check("summary transactions count", wb["Summary"].cell(4, 2).value == 5,
                str(wb["Summary"].cell(4, 2).value))

    r = requests.post(
        BASE + "/convert",
        files={"statement": ("x.txt", b"hello", "text/plain")},
        data={"password": ""},
        allow_redirects=False,
    )
    ok &= check("txt rejected", r.status_code == 302, str(r.status_code))

    r = requests.post(
        BASE + "/convert",
        files={"statement": ("big.pdf", b"0" * (21 * 1024 * 1024), "application/pdf")},
        data={"password": ""},
        allow_redirects=False,
    )
    ok &= check("oversize rejected", r.status_code in (302, 413), str(r.status_code))

    r = requests.post(BASE + "/convert", data={}, allow_redirects=False)
    ok &= check("missing file rejected", r.status_code == 302, str(r.status_code))

    print("\nALL PASSED" if ok else "\nSOME CHECKS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
